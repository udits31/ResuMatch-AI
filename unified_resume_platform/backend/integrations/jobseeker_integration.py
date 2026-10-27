import sys
import os

sys.path.append(os.path.join(os.path.dirname(__file__), '..', '..'))

try:
    from ..models.job_analyzer import JobAnalyzer
    from ..models.profile_manager import ProfileManager
    from ..models.resume_generator import ResumeGenerator
    from ..models.matching_engine import ResumeMatcher
    from ..database.db_manager import DatabaseManager
    from ..database.profile_repository import ProfileRepository
except ImportError as e:
    print(f"Error importing Job Seeker modules: {e}")
    print(f"Current working directory: {os.getcwd()}")
    print(f"Python path: {sys.path}")
    raise

class JobSeekerIntegration:
    def __init__(self):
        # Set correct data directory path for job seeker models
        data_dir = os.path.join(os.path.dirname(__file__), '..', '..', 'data')
        
        self.job_analyzer = JobAnalyzer()
        self.profile_manager = ProfileManager(data_dir)
        self.resume_generator = ResumeGenerator()
        self.matcher = ResumeMatcher()
        self.db_manager = DatabaseManager()
        self.profile_repo = ProfileRepository()
        self.current_analysis = None
        self.current_profile = None

    def analyze_job(self, job_description):
        try:
            if not job_description.strip():
                return {
                    'success': False,
                    'message': 'Job description is required',
                    'data': None,
                    'errors': ['Empty job description']
                }

            keywords = self.job_analyzer.extract_keywords(job_description)
            requirements = self.job_analyzer.identify_requirements(job_description)

            analysis = {
                'keywords': keywords,
                'relevance_score': 0.0,
                **requirements
            }

            # Fit analysis against the saved profile (if one exists)
            profile = self.current_profile or self.profile_repo.load()
            if profile:
                fit = self._compute_fit(profile, job_description)
                analysis['relevance_score'] = round(fit['overall'] / 100.0, 4)
                analysis['fit'] = fit

            self.current_analysis = analysis

            # Database saving is optional for job analysis
            # Removed to keep the system simple and focused

            return {
                'success': True,
                'message': 'Job analysis completed successfully',
                'data': analysis,
                'errors': []
            }
        except Exception as e:
            return {
                'success': False,
                'message': 'Error analyzing job description',
                'data': None,
                'errors': [str(e)]
            }

    def _profile_to_text(self, profile):
        """Flatten a structured profile into a single resume-like text blob for matching."""
        parts = []
        if profile.get('summary'):
            parts.append(profile['summary'])
        for exp in profile.get('work_experience', []) or []:
            parts.append(f"{exp.get('position', '')} at {exp.get('company', '')}")
            if exp.get('description'):
                parts.append(exp['description'])
            parts.extend(exp.get('achievements', []) or [])
        for edu in profile.get('education', []) or []:
            parts.append(f"{edu.get('degree', '')} {edu.get('field', '')} {edu.get('institution', '')}")
        skills = profile.get('skills', {}) or {}
        for category in ['technical', 'soft', 'languages', 'certifications']:
            values = skills.get(category)
            if isinstance(values, list):
                parts.extend(values)
        for proj in profile.get('projects', []) or []:
            parts.append(proj.get('name') or proj.get('title', ''))
            if proj.get('description'):
                parts.append(proj['description'])
            parts.extend(proj.get('technologies', []) or [])
        for pub in profile.get('publications', []) or []:
            parts.append(pub.get('title', ''))
            if pub.get('description'):
                parts.append(pub['description'])
        for act in profile.get('extracurricular', []) or []:
            parts.append(act.get('title', ''))
            if act.get('description'):
                parts.append(act['description'])
        return "\n".join(p for p in parts if p)

    def _compute_fit(self, profile, job_description):
        """Score the profile against the job description and build tailoring suggestions."""
        profile_text = self._profile_to_text(profile)
        result = self.matcher.get_match_analysis(profile_text, job_description)
        detected_skills = self.matcher.extract_skills(profile_text)

        overall = result['match_score']
        if overall >= 70:
            verdict = 'Strong fit'
        elif overall >= 40:
            verdict = 'Moderate fit'
        else:
            verdict = 'Low fit'

        return {
            'overall': overall,
            'semantic': result['semantic_score'],
            'skill_match': result['skill_match_percentage'],
            'matching_skills': result['matching_skills'],
            'missing_skills': result['missing_skills'],
            'detected_skills': detected_skills,
            'verdict': verdict,
            'suggestions': self._build_suggestions(result, profile),
        }

    def _build_suggestions(self, result, profile):
        """Turn the skill gap and profile state into actionable tailoring tips."""
        suggestions = []
        missing = result.get('missing_skills', [])
        matching = result.get('matching_skills', [])

        if missing:
            suggestions.append(
                "Add these job-relevant skills if you have them: " + ", ".join(missing[:10])
            )
        if result.get('skill_match_percentage', 0) < 50 and result.get('jd_skills'):
            suggestions.append(
                "Your skill overlap with this job is low. Mirror the job's wording in your experience bullets where it is true for you."
            )
        if not (profile.get('summary') or '').strip():
            suggestions.append(
                "Add a professional summary highlighting your most relevant experience for this role."
            )
        if matching:
            suggestions.append(
                "Lead with your matching strengths: " + ", ".join(matching[:8])
            )
        if not suggestions:
            suggestions.append("Strong alignment. You look like a great match for this role.")
        return suggestions

    def get_profile(self, name=None):
        """Get a profile by name, or the most recently updated one."""
        try:
            profile = self.profile_repo.load(name)

            if profile:
                self.current_profile = profile
                return {
                    'success': True,
                    'message': 'Profile loaded successfully',
                    'data': profile,
                    'errors': []
                }
            else:
                return {
                    'success': True,
                    'message': 'No profile found',
                    'data': None,
                    'errors': []
                }
        except Exception as e:
            return {
                'success': False,
                'message': 'Error loading profile',
                'data': None,
                'errors': [str(e)]
            }
    
    def list_profiles(self):
        """Return the names of all saved profiles."""
        try:
            return {
                'success': True,
                'message': 'Profiles listed',
                'data': {'profiles': self.profile_repo.list_names()},
                'errors': []
            }
        except Exception as e:
            return {'success': False, 'message': 'Error listing profiles', 'data': None, 'errors': [str(e)]}

    def create_profile(self, profile_data):
        """Create a new profile"""
        try:
            # Validate profile data
            errors = self.profile_manager.validate_profile(profile_data)
            if errors:
                return {
                    'success': False,
                    'message': 'Profile validation failed',
                    'data': None,
                    'errors': errors
                }
            
            name = (profile_data.get('profile_name')
                    or (profile_data.get('personal_info', {}) or {}).get('name')
                    or 'Default')
            profile_data['profile_name'] = name
            success = self.profile_repo.save(name, profile_data)
            if success:
                self.current_profile = profile_data
                return {
                    'success': True,
                    'message': 'Profile created successfully',
                    'data': profile_data,
                    'errors': []
                }
            else:
                return {
                    'success': False,
                    'message': 'Failed to create profile',
                    'data': None,
                    'errors': ['Profile creation failed']
                }
        except Exception as e:
            return {
                'success': False,
                'message': 'Error creating profile',
                'data': None,
                'errors': [str(e)]
            }
    
    def update_profile(self, profile_data):
        """Update existing profile"""
        try:
            # Validate profile data
            errors = self.profile_manager.validate_profile(profile_data)
            if errors:
                return {
                    'success': False,
                    'message': 'Profile validation failed',
                    'data': None,
                    'errors': errors
                }
            
            name = (profile_data.get('profile_name')
                    or (profile_data.get('personal_info', {}) or {}).get('name')
                    or 'Default')
            profile_data['profile_name'] = name
            success = self.profile_repo.save(name, profile_data)
            if success:
                self.current_profile = profile_data
                return {
                    'success': True,
                    'message': 'Profile updated successfully',
                    'data': profile_data,
                    'errors': []
                }
            else:
                return {
                    'success': False,
                    'message': 'Failed to update profile',
                    'data': None,
                    'errors': ['Profile update failed']
                }
        except Exception as e:
            return {
                'success': False,
                'message': 'Error updating profile',
                'data': None,
                'errors': [str(e)]
            }
    
    def validate_profile(self, profile_data):
        """Validate profile data"""
        try:
            errors = self.profile_manager.validate_profile(profile_data)
            return {
                'success': len(errors) == 0,
                'message': 'Profile is valid' if len(errors) == 0 else 'Profile validation failed',
                'data': {'valid': len(errors) == 0, 'errors': errors},
                'errors': errors
            }
        except Exception as e:
            return {
                'success': False,
                'message': 'Error validating profile',
                'data': None,
                'errors': [str(e)]
            }
    
    def generate_resume(self, profile=None, analysis=None):
        """Generate tailored resume"""
        try:
            if not profile:
                profile = self.current_profile
            if not analysis:
                analysis = self.current_analysis
            
            if not profile:
                return {
                    'success': False,
                    'message': 'Profile is required for resume generation',
                    'data': None,
                    'errors': ['No profile available']
                }
            
            if not analysis:
                return {
                    'success': False,
                    'message': 'Job analysis is required for resume generation',
                    'data': None,
                    'errors': ['No job analysis available']
                }
            
            resume_data = self.resume_generator.generate_resume(profile, analysis)
            formatted_resume = self.resume_generator.format_resume(resume_data)
            
            return {
                'success': True,
                'message': 'Resume generated successfully',
                'data': {
                    'resume_data': resume_data,
                    'formatted_resume': formatted_resume
                },
                'errors': []
            }
        except Exception as e:
            return {
                'success': False,
                'message': 'Error generating resume',
                'data': None,
                'errors': [str(e)]
            }
    
    def export_resume(self, content, format_type, profile=None, job_description=None, resume_data=None):
        """Export resume in specified format for web download"""
        try:
            import tempfile
            import os
            from datetime import datetime
            
            if not content:
                return {
                    'success': False,
                    'message': 'No resume content available for export',
                    'data': None,
                    'errors': ['No resume content']
                }
            
            # Generate filename
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            default_name = "resume"
            if profile and "personal_info" in profile:
                name = profile["personal_info"].get("name", "").strip()
                if name:
                    default_name = "".join(c for c in name if c.isalnum() or c in (' ', '-', '_')).strip()
                    default_name = default_name.replace(' ', '_').lower()
            
            filename = f"{default_name}_{timestamp}.{format_type}"
            
            # Generate file content based on format
            if format_type == 'txt':
                file_data = content.encode('utf-8')
                content_type = 'text/plain'
            elif format_type == 'pdf':
                if resume_data:
                    try:
                        docx_bytes = self._build_formatted_docx(resume_data)
                        file_data = self._docx_to_pdf_bytes(docx_bytes)
                    except Exception as conv_err:
                        print(f"Word PDF conversion failed, using fallback PDF: {conv_err}")
                        file_data = self._generate_pdf(content)
                else:
                    file_data = self._generate_pdf(content)
                content_type = 'application/pdf'
            elif format_type == 'docx':
                if resume_data:
                    file_data = self._build_formatted_docx(resume_data)
                else:
                    file_data = self._generate_docx(content)
                content_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            else:
                return {
                    'success': False,
                    'message': 'Unsupported export format',
                    'data': None,
                    'errors': ['Format must be txt, pdf, or docx']
                }
            
            return {
                'success': True,
                'message': 'Resume exported successfully',
                'data': {
                    'filename': filename,
                    'file_data': file_data,
                    'content_type': content_type,
                    'format': format_type
                },
                'errors': []
            }
        except Exception as e:
            return {
                'success': False,
                'message': 'Error exporting resume',
                'data': None,
                'errors': [str(e)]
            }
    
    def _generate_pdf(self, content):
        """Generate PDF content"""
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from io import BytesIO
            
            buffer = BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=letter)
            styles = getSampleStyleSheet()
            story = []
            
            # Split content into paragraphs
            paragraphs = content.split('\n\n')
            for para in paragraphs:
                if para.strip():
                    # Check if it's a header (all caps or starts with specific patterns)
                    if para.isupper() or para.startswith('=') or para.startswith('-'):
                        story.append(Paragraph(para.strip(), styles['Heading2']))
                    else:
                        story.append(Paragraph(para.strip(), styles['Normal']))
                    story.append(Spacer(1, 12))
            
            doc.build(story)
            buffer.seek(0)
            return buffer.getvalue()
        except ImportError:
            # Fallback to text if reportlab not available
            return content.encode('utf-8')
        except Exception as e:
            print(f"PDF generation error: {e}")
            return content.encode('utf-8')
    
    def _generate_docx(self, content):
        """Generate DOCX content"""
        try:
            from docx import Document
            from io import BytesIO
            
            doc = Document()
            
            # Split content into paragraphs
            paragraphs = content.split('\n\n')
            for para in paragraphs:
                para = para.strip()
                if not para:
                    continue
                
                # Check if it's a header
                if para.isupper() or para.startswith('=') or para.startswith('-'):
                    doc.add_heading(para.replace('=', '').replace('-', '').strip(), level=2)
                else:
                    doc.add_paragraph(para)
            
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        except ImportError:
            # Fallback to text if python-docx not available
            return content.encode('utf-8')
        except Exception as e:
            print(f"DOCX generation error: {e}")
            return content.encode('utf-8')

    def _build_formatted_docx(self, resume_data):
        """Build a professionally formatted resume .docx from structured resume_data:
        centered name/contact header, bold underlined section headings, bold entry
        titles with right-aligned dates, and bulleted achievements."""
        from io import BytesIO
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        RIGHT_TAB = Inches(7.3)
        ACCENT = RGBColor(0x2F, 0x4F, 0x4F)

        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        normal = doc.styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(2)

        def bottom_border(paragraph):
            pPr = paragraph._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '2')
            bottom.set(qn('w:color'), '2F4F4F')
            pbdr.append(bottom)
            pPr.append(pbdr)

        def add_heading(text):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text.upper())
            r.bold = True
            r.font.color.rgb = ACCENT
            bottom_border(p)

        def add_title_with_date(title, date=''):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
            rt = p.add_run(title)
            rt.bold = True
            if date:
                p.add_run('\t' + date)

        def add_bullet(text):
            doc.add_paragraph(text, style='List Bullet')

        # Header: name + contact (centered)
        pi = resume_data.get('personal_info', {}) or {}
        if pi.get('name'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(pi['name'].upper())
            r.bold = True
            r.font.size = Pt(16)
        if pi.get('headline'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(pi['headline'])
            r.font.size = Pt(11)
            r.font.color.rgb = ACCENT
        contact = [pi.get(f) for f in ['email', 'phone', 'location', 'address', 'linkedin'] if pi.get(f)]
        if contact:
            p = doc.add_paragraph(' | '.join(contact))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.size = Pt(10)

        # Summary
        if resume_data.get('summary'):
            add_heading('Professional Summary')
            doc.add_paragraph(resume_data['summary'])

        # Education
        education = resume_data.get('education', []) or []
        if education:
            add_heading('Education')
            for edu in education:
                add_title_with_date(edu.get('institution', ''), edu.get('graduation_date', ''))
                degree = edu.get('degree', '')
                if edu.get('field'):
                    degree += (' in ' if degree else '') + edu['field']
                if edu.get('gpa'):
                    degree += (' | ' if degree else '') + f"GPA: {edu['gpa']}"
                if degree:
                    doc.add_paragraph(degree)
                    
        # Work experience
        work_experience = resume_data.get('work_experience', []) or []
        if work_experience:
            add_heading('Work Experience')
        for exp in work_experience:
            title = exp.get('position', '')
            if exp.get('company'):
                title += (' | ' if title else '') + exp['company']
            date = ''
            if exp.get('start_date'):
                date = exp['start_date'] + ' - ' + (exp.get('end_date') or 'Present')
            add_title_with_date(title, date)
            if exp.get('description'):
                add_bullet(exp['description'])
            for achievement in (exp.get('achievements') or [])[:4]:
                add_bullet(achievement)

        # Projects
        projects = resume_data.get('projects', []) or []
        if projects:
            add_heading('Projects')
            for proj in projects:
                title = proj.get('name') or proj.get('title') or ''
                if proj.get('technologies'):
                    title += (' | ' if title else '') + ', '.join(proj['technologies'][:6])
                add_title_with_date(title)
                if proj.get('description'):
                    add_bullet(proj['description'])
                link = proj.get('url') or proj.get('github')
                if link:
                    doc.add_paragraph(link)
        # Publications
        publications = resume_data.get('publications', []) or []
        if publications:
            add_heading('Publications')
            for pub in publications:
                title = pub.get('title', '')
                if pub.get('venue'):
                    title += (' | ' if title else '') + pub['venue']
                add_title_with_date(title)
                if pub.get('description'):
                    add_bullet(pub['description'])
                if pub.get('url'):
                    doc.add_paragraph(pub['url'])
                    
        # Skills
        skills = resume_data.get('skills', {}) or {}
        skill_labels = {
            'technical': 'Technical',
            'soft': 'Soft Skills',
            'languages': 'Languages',
            'certifications': 'Certifications',
        }
        if any(skills.get(k) for k in skill_labels):
            add_heading('Skills')
            for key, label in skill_labels.items():
                values = skills.get(key)
                if values:
                    p = doc.add_paragraph(style='List Bullet')
                    r = p.add_run(label + ': ')
                    r.bold = True
                    p.add_run(', '.join(values[:12]))

        # Extra-curricular / Leadership
        extracurricular = resume_data.get('extracurricular', []) or []
        if extracurricular:
            add_heading('Extra-curricular / Leadership')
            for act in extracurricular:
                title = act.get('title', '')
                if act.get('organization'):
                    title += (' | ' if title else '') + act['organization']
                add_title_with_date(title)
                if act.get('description'):
                    add_bullet(act['description'])

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _docx_to_pdf_bytes(self, docx_bytes):
        """Convert a .docx (bytes) to PDF (bytes) using MS Word via docx2pdf.
        Initializes COM for the current thread (Flask serves requests off the
        main thread) and cleans up temp files."""
        import os
        import tempfile
        from docx2pdf import convert

        co_module = None
        try:
            import comtypes
            comtypes.CoInitialize()
            co_module = comtypes
        except Exception:
            try:
                import pythoncom
                pythoncom.CoInitialize()
                co_module = pythoncom
            except Exception:
                co_module = None

        tmp_dir = tempfile.mkdtemp()
        docx_path = os.path.join(tmp_dir, 'resume.docx')
        pdf_path = os.path.join(tmp_dir, 'resume.pdf')
        try:
            with open(docx_path, 'wb') as f:
                f.write(docx_bytes)
            convert(docx_path, pdf_path)
            with open(pdf_path, 'rb') as f:
                return f.read()
        finally:
            try:
                if co_module:
                    co_module.CoUninitialize()
            except Exception:
                pass
            for path in (docx_path, pdf_path):
                try:
                    os.remove(path)
                except Exception:
                    pass
            try:
                os.rmdir(tmp_dir)
            except Exception:
                pass